"""Builds the IEEE-style Word version of the paper from main.tex + refs.bib, so both stay identical.

Handles only the LaTeX subset used in main.tex (sections, emphasis, citations, lists, booktabs
tables, figures, simple inline math). Output: paper/Brain_Age_Audit_IEEE.docx
Usage: python paper/build_docx.py
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEX = (HERE / "main.tex").read_text(encoding="utf8")
BIB = (HERE / "refs.bib").read_text(encoding="utf8")
FONT = "Times New Roman"
COL_W, PAGE_W = 3.5, 7.16

# --------------------------------------------------------------------------------------------
# Bibliography: parse, order by first citation, format in IEEE style
# --------------------------------------------------------------------------------------------

def parse_bib(text):
    entries = {}
    for m in re.finditer(r"@(\w+)\{([^,]+),(.*?)\}\s*(?=@|\Z)", text, re.S):
        kind, key, body = m.group(1).lower(), m.group(2).strip(), m.group(3)
        fields = {}
        for f in re.finditer(r"(\w+)\s*=\s*\{((?:[^{}]|\{(?:[^{}]|\{[^{}]*\})*\})*)\}", body):
            fields[f.group(1).lower()] = f.group(2)
        entries[key] = (kind, fields)
    return entries


BIBS = parse_bib(BIB)
ORDER = []
for m in re.finditer(r"\\cite\{([^}]*)\}", TEX):
    for k in m.group(1).split(","):
        k = k.strip()
        if k not in ORDER:
            ORDER.append(k)
NUM = {k: i + 1 for i, k in enumerate(ORDER)}


def tex_chars(s):
    s = s.replace("{\\c{c}}", "ç").replace("\\c{c}", "ç").replace('{\\"o}', "ö").replace('{\\"u}', "ü")
    s = s.replace("{\\'a}", "á").replace("{\\'e}", "é").replace("{\\'i}", "í").replace("{\\'u}", "ú")
    s = s.replace("\\&", "&").replace("--", "–").replace("\\,", " ")
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    return s.replace("{", "").replace("}", "")


def ieee_authors(a):
    names = [n.strip() for n in a.split(" and ")]
    out = []
    for n in names:
        if n == "others":
            out.append("et al.")
            continue
        last, _, first = n.partition(",")
        initials = " ".join(p[0] + "." for p in re.split(r"[\s]+", first.strip()) if p and p[0].isalpha())
        initials = initials.replace("..", ".")
        out.append(f"{initials} {last.strip()}".strip())
    if out and out[-1] == "et al.":
        return ", ".join(out[:-1]) + " et al."
    if len(out) <= 2:
        return " and ".join(out)
    return ", ".join(out[:-1]) + ", and " + out[-1]


def ref_runs(key):
    """List of (text, italic) runs for one reference."""
    kind, f = BIBS[key]
    g = lambda k: tex_chars(f.get(k, ""))
    if kind == "misc":
        return [(f"{g('title')}. [Online]. Available: {g('howpublished')}", False)]
    runs = [(f"{ieee_authors(tex_chars(f['author']))}, “{g('title')},” ", False)]
    if kind == "article":
        j = g("journal")
        if j.startswith("arXiv"):
            return runs + [(f"{j}, {g('year')}.", False)]
        runs.append((j, True))
        tail = ""
        if f.get("volume"): tail += f", vol. {g('volume')}"
        if f.get("number"): tail += f", no. {g('number')}"
        if f.get("pages"): tail += f", {'pp.' if '–' in g('pages') else 'Art. no.'} {g('pages')}"
        return runs + [(f"{tail}, {g('year')}.", False)]
    if kind == "inproceedings":
        runs += [("in ", False), (g("booktitle"), True)]
        tail = f", vol. {g('volume')}" if f.get("volume") else ""
        tail += f", pp. {g('pages')}" if f.get("pages") else ""
        return runs + [(f"{tail}, {g('year')}.", False)]
    if kind == "book":
        return [(f"{ieee_authors(tex_chars(f['author']))}, ", False), (g("title"), True),
                (f". {g('address')}: {g('publisher')}, {g('year')}.", False)]
    raise ValueError(kind)


def cite_text(keys):
    nums = sorted(NUM[k.strip()] for k in keys.split(","))
    groups, start = [], nums[0]
    for a, b in zip(nums, nums[1:] + [None]):
        if b != a + 1:
            groups.append(f"[{start}]" if start == a else (f"[{start}], [{a}]" if a == start + 1 else f"[{start}]–[{a}]"))
            start = b
    return ", ".join(groups)

# --------------------------------------------------------------------------------------------
# Inline LaTeX -> runs
# --------------------------------------------------------------------------------------------

MATH = [(r"\sum_k p_k c_k", "Σₖ pₖcₖ"), (r"\hat{y}", "ŷ"), (r"\Delta", "Δ"), (r"\times", "×"),
        (r"\rightarrow", "→"), (r"\sigma", "σ"), (r"\rho", "ρ"), (r"\gamma", "γ"), (r"\sim", " ~ "),
        (r"\mathrm{BAG}", "BAG"), (r"\mathrm{sex}", "sex"), (r"\mathrm{site}", "site"),
        (r"\mathrm{age}", "age"), (r"\mathrm{share}", "share"), (r"\mathrm{error}", "error"),
        ("|", "|"), ("^{-3}", "⁻³"), ("^{-4}", "⁻⁴"), ("^{-11}", "⁻¹¹"), ("^2", "²"), (r"\,", " "),
        (r"\cdot", "·")]


SUP = str.maketrans("0123456789-+", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺")


def math(s):
    s = s.replace(r"\pm", "±")
    s = re.sub(r"\^\{([-+]?\d+)\}", lambda m: m.group(1).translate(SUP), s)
    bold = s.startswith(r"\mathbf{")
    if bold:
        s = s[len(r"\mathbf{"):-1]
    for a, b in MATH:
        s = s.replace(a, b)
    s = s.replace("{", "").replace("}", "").replace("-", "−").replace("+", "+")
    return s, bold


def clean(s):
    s = re.sub(r"\\label\{[^}]*\}", "", s)
    s = re.sub(r"\\textsc\{([^}]*)\}", lambda m: m.group(1).upper(), s)
    s = re.sub(r"\\pend\{([^}]*)\}", lambda m: "[" + (m.group(1) or "pending") + "]", s)
    s = re.sub(r"\\ci\{([^}]*)\}\{([^}]*)\}", lambda m: "[" + m.group(1).replace("-", "−") + ", " + m.group(2).replace("-", "−") + "]", s)
    s = s.replace("~", " ").replace("\\%", "%").replace("---", "—").replace("--", "–")
    s = s.replace("``", "“").replace("''", "”").replace("\\&", "&").replace("\\_", "_")
    s = s.replace("\\,", " ").replace("Pi\\c{c}arra", "Piçarra").replace("\\\\", " ")
    return s


TOKEN = re.compile(r"\\(textbf|emph|textit|texttt|url|textcolor\{[^}]*\})\{|\\cite\{([^}]*)\}|"
                   r"\\ref\{([^}]*)\}|\$([^$]*)\$|\}")


def runs(s, bold=False, italic=False):
    """Parse inline LaTeX into (text, bold, italic, mono) runs; nesting handled with a stack."""
    s = clean(s)
    out, stack, pos = [], [], 0
    state = lambda: (bold or any(x == "b" for x in stack), italic or any(x == "i" for x in stack),
                     any(x == "t" for x in stack))
    for m in TOKEN.finditer(s):
        if m.start() > pos:
            out.append((s[pos:m.start()], *state()))
        pos = m.end()
        if m.group(1):
            c = m.group(1)
            stack.append("b" if c == "textbf" else "i" if c in ("emph", "textit") else "t" if c in ("texttt", "url") else "n")
        elif m.group(2) is not None:
            out.append((cite_text(m.group(2)), *state()))
        elif m.group(3) is not None:
            out.append((LABELS.get(m.group(3), "?"), *state()))
        elif m.group(4) is not None:
            t, b = math(m.group(4))
            st = state()
            out.append((t, st[0] or b, st[1], False))
        else:
            if stack:
                stack.pop()
    if pos < len(s):
        out.append((s[pos:], *state()))
    return [(t.replace("{", "").replace("}", ""), b, i, mono) for t, b, i, mono in out if t]

# --------------------------------------------------------------------------------------------
# Document helpers
# --------------------------------------------------------------------------------------------

def set_cols(section, n):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")


def para(doc, text="", size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold=False, italic=False,
         before=0, after=0, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment, pf.space_before, pf.space_after = align, Pt(before), Pt(after)
    pf.first_line_indent = Inches(0.17) if indent else None
    pf.line_spacing = 1.0
    pf.keep_with_next = keep
    for t, b, i, mono in (runs(text, bold, italic) if isinstance(text, str) else text):
        r = p.add_run(t)
        r.bold, r.italic = b, i
        r.font.size = Pt(size)
        r.font.name = "Consolas" if mono else FONT
    return p


def roman(n):
    return ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"][n]


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)


def borders(table, rows_with_top, n_rows):
    """Booktabs look: top rule, rule under header, bottom rule."""
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            for side in ("top", "bottom"):
                el = OxmlElement(f"w:{side}")
                on = (side == "top" and ri in rows_with_top) or (side == "bottom" and ri == n_rows - 1)
                el.set(qn("w:val"), "single" if on else "nil")
                el.set(qn("w:sz"), "8" if on else "0")
                b.append(el)
            for side in ("left", "right"):
                el = OxmlElement(f"w:{side}"); el.set(qn("w:val"), "nil"); b.append(el)
            tcPr.append(b)

# --------------------------------------------------------------------------------------------
# Pre-pass: number figures and tables so \ref works
# --------------------------------------------------------------------------------------------

body = TEX[TEX.index(r"\end{IEEEkeywords}") + len(r"\end{IEEEkeywords}"):TEX.index(r"\bibliographystyle")]
LABELS, nfig, ntab = {}, 0, 0
_ns = _nsub = 0
for m in re.finditer(r"\\(section|subsection)\{[^}]*\}\s*(\\label\{([^}]*)\})?", body):
    if m.group(1) == "section":
        _ns += 1; _nsub = 0
    else:
        _nsub += 1
    if m.group(3):
        LABELS[m.group(3)] = roman(_ns) + ("-" + chr(64 + _nsub) if m.group(1) == "subsection" else "")
for m in re.finditer(r"\\begin\{(figure\*?|table)\}.*?\\end\{\1\}", body, re.S):
    lab = re.search(r"\\label\{([^}]*)\}", m.group(0)).group(1)
    if m.group(1).startswith("figure"):
        nfig += 1; LABELS[lab] = str(nfig)
    else:
        ntab += 1; LABELS[lab] = roman(ntab)

# --------------------------------------------------------------------------------------------
# Build
# --------------------------------------------------------------------------------------------

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.67)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

title = re.search(r"\\title\{(.*?)\}\n", TEX).group(1)
para(doc, title, size=24, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=10)

authors = [("Adharshan N*", "School of Computer Science and Engineering", "adharshan2706@gmail.com"),
           ("Jaaiwanth K*", "School of Computer Science and Engineering", "jaaiwanth90@gmail.com"),
           ("Shridevi S†", "Centre for Neuroinformatics", "[faculty email]")]
tbl = doc.add_table(rows=1, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, (name, dept, mail) in zip(tbl.rows[0].cells, authors):
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell.paragraphs[0].add_run(name); r.font.size = Pt(11)
    for line, it in [(dept, True), ("Vellore Institute of Technology", True), ("Chennai, India", False), (mail, False)]:
        p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); r.italic = it; r.font.size = Pt(9)
para(doc, "*Equal contribution.  †Corresponding author.", size=8, align=WD_ALIGN_PARAGRAPH.CENTER,
     indent=False, italic=True, before=4, after=8)

sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
set_cols(sec2, 2)

abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", TEX, re.S).group(1).strip()
para(doc, [("Abstract", True, True, False), ("—" , True, False, False)] +
     [(t, True, i, m) for t, b, i, m in runs(abstract)], size=9, indent=False, after=6)
kw = re.search(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", TEX, re.S).group(1).strip()
para(doc, [("Index Terms", True, True, False), ("—" + kw[0].upper() + kw[1:] + ".", True, False, False)],
     size=9, indent=False, after=6)

n_sec = n_sub = 0
blocks = re.split(r"(\\section\*?\{[^}]*\}|\\subsection\{[^}]*\}|\\begin\{(?:figure\*?|table|itemize|enumerate)\}.*?\\end\{(?:figure\*?|table|itemize|enumerate)\})",
                  body, flags=re.S)


def add_figure(blk, wide):
    img = re.search(r"\\includegraphics\[width=([\d.]*)\\(textwidth|columnwidth)\]\{([^}]*)\}", blk)
    imgs = re.findall(r"\\includegraphics\[width=([\d.]*)\\(textwidth|columnwidth)\]\{([^}]*)\}", blk)
    cap = re.search(r"\\caption\{(.*)\}\s*\n\\label", blk, re.S).group(1)
    lab = re.search(r"\\label\{([^}]*)\}", blk).group(1)
    if wide:
        s = doc.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 1)
    for frac, unit, name in imgs:
        f = HERE / "figures" / name
        if not f.exists() or f.suffix == ".pdf":
            f = f.with_suffix(".png")
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        width = (PAGE_W if unit == "textwidth" else COL_W) * (float(frac) if frac else 1.0)
        p.add_run().add_picture(str(f), width=Inches(width))
    para(doc, [(f"Fig. {LABELS[lab]}.  ", False, False, False)] + runs(cap), size=8,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY if not wide else WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=8)
    if wide:
        s = doc.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 2)


def add_table(blk):
    cap = re.search(r"\\caption\{(.*?)\}\n", blk).group(1)
    lab = re.search(r"\\label\{([^}]*)\}", blk).group(1)
    para(doc, f"TABLE {LABELS[lab]}", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6, keep=True)
    para(doc, cap.upper() if len(cap) < 0 else cap, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
         after=3, keep=True)
    tab = re.search(r"\\begin\{tabular\}\{(?:[^{}]|\{[^{}]*\})*\}(.*?)\\end\{tabular\}", blk, re.S).group(1)
    note = None
    rows = []
    for line in tab.split("\\\\"):
        line = re.sub(r"\\(toprule|midrule|bottomrule)", "", line).strip()
        if not line:
            continue
        mc = re.match(r"\\multicolumn\{\d+\}\{[^}]*\}\{(.*)\}$", line, re.S)
        if mc:
            note = re.sub(r"\\footnotesize\s*", "", mc.group(1)); continue
        rows.append([c.strip() for c in re.split(r"(?<!\\)&", line)])
    mids = [i for i, l in enumerate([l for l in tab.split("\\\\") if re.sub(r"\\(toprule|midrule|bottomrule)", "", l).strip() and "multicolumn" not in l])]
    # rows that start after a \midrule get a top border
    rule_rows, k = {0, 1}, 0
    for line in tab.split("\\\\"):
        stripped = re.sub(r"\\(toprule|midrule|bottomrule)", "", line).strip()
        if "\\midrule" in line and stripped == "" :
            rule_rows.add(k)
        elif "\\midrule" in line:
            rule_rows.add(k)
        if stripped and "multicolumn" not in stripped:
            k += 1
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
            for txt, b, i, m in runs(cell_text, bold=False):
                r = p.add_run(txt); r.bold, r.italic = b or ri == 0, i; r.font.size = Pt(8); r.font.name = FONT
                if "PASS" in txt: r.font.color.rgb = RGBColor(0x0F, 0x7A, 0x3A)
                if "FAIL" in txt: r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
    borders(t, rule_rows, len(rows))
    if note:
        para(doc, note, size=7.5, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
    para(doc, "", size=4)


def add_list(blk, ordered):
    items = re.split(r"\\item\s*", re.sub(r"\\(begin|end)\{(itemize|enumerate)\}", "", blk))[1:]
    for n, it in enumerate(items, 1):
        p = para(doc, [(f"{n})  " if ordered else "•  ", False, False, False)] + runs(it.strip()),
                 indent=False, after=1)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.15)


for blk in blocks:
    if not blk or not blk.strip():
        continue
    if blk.startswith("\\section*"):
        name = re.search(r"\{(.*)\}", blk).group(1)
        para(doc, name.upper(), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=8, after=4, keep=True)
    elif blk.startswith("\\section"):
        n_sec += 1; n_sub = 0
        name = re.search(r"\{(.*)\}", blk).group(1)
        para(doc, f"{roman(n_sec)}. {name.upper()}", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
             before=8, after=4, keep=True)
    elif blk.startswith("\\subsection"):
        n_sub += 1
        name = re.search(r"\{(.*)\}", blk).group(1)
        para(doc, f"{chr(64 + n_sub)}. {name}", italic=True, indent=False, before=4, after=2, keep=True,
             align=WD_ALIGN_PARAGRAPH.LEFT)
    elif blk.startswith("\\begin{figure*}"):
        add_figure(blk, True)
    elif blk.startswith("\\begin{figure}"):
        add_figure(blk, False)
    elif blk.startswith("\\begin{table}"):
        add_table(blk)
    elif blk.startswith("\\begin{itemize}"):
        add_list(blk, False)
    elif blk.startswith("\\begin{enumerate}"):
        add_list(blk, True)
    else:
        text = re.sub(r"(?m)^%.*$", "", blk)
        for chunk in re.split(r"\n\s*\n", text):
            chunk = " ".join(chunk.split())
            if chunk:
                para(doc, chunk, after=2)

para(doc, "REFERENCES", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=8, after=4, keep=True)
for k in ORDER:
    p = para(doc, [(f"[{NUM[k]}]\t", False, False, False)] + [(t, False, i, False) for t, i in ref_runs(k)],
             size=8, indent=False, after=1, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))

out = HERE / "Brain_Age_Audit_IEEE.docx"
doc.save(out)
print("saved", out, "| refs:", len(ORDER), "| figs:", nfig, "| tables:", ntab)
missing = [k for k in ORDER if k not in BIBS]
print("missing bib keys:", missing)
print("unused bib keys:", [k for k in BIBS if k not in NUM])
